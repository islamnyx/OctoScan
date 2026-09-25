"""Client report conversion — Markdown (constrained subset) to DOCX / PDF.

The Report Agent prompt restricts output to: ``#``/``##``/``###`` headings,
paragraphs, ``- `` / ``1. `` lists, fenced code blocks, ``**bold**``, and
GitHub tables.  This module parses exactly that subset into blocks and
renders two thin backends (python-docx, reportlab platypus).  Unknown lines
degrade to plain paragraphs — the converter never crashes on model output.

Backends are imported lazily so a broken install fails with a clear message
instead of an ImportError at CLI startup.
"""

from __future__ import annotations

import html
import logging
import re
from pathlib import Path

log = logging.getLogger(__name__)

SUPPORTED_FORMATS = ("md", "docx", "pdf")

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


# ---------------------------------------------------------------------------
# Block model + parser
# ---------------------------------------------------------------------------

# Block = (kind, payload): kinds h1/h2/h3/para/ulist/olist/code/table.

def parse_markdown(text: str) -> list[tuple[str, object]]:
    blocks: list[tuple[str, object]] = []
    para_lines: list[str] = []
    list_items: list[str] = []
    list_kind: str | None = None
    code_lines: list[str] = []
    in_code = False
    table_rows: list[list[str]] = []

    def flush_para():
        if para_lines:
            blocks.append(("para", " ".join(para_lines)))
            para_lines.clear()

    def flush_list():
        nonlocal list_kind
        if list_items:
            blocks.append((list_kind or "ulist", list(list_items)))
            list_items.clear()
            list_kind = None

    def flush_table():
        if table_rows:
            header, *rows = table_rows
            blocks.append(("table", (header, rows)))
            table_rows.clear()

    def is_table_sep(cells: list[str]) -> bool:
        return bool(cells) and all(
            re.fullmatch(r":?-{2,}:?", c.strip()) for c in cells)

    for raw in (text or "").splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            if in_code:
                blocks.append(("code", list(code_lines)))
                code_lines.clear()
                in_code = False
            else:
                flush_para()
                flush_list()
                flush_table()
                in_code = True
            continue
        if in_code:
            code_lines.append(raw.rstrip("\n"))
            continue
        stripped = line.strip()
        if not stripped:
            flush_para()
            flush_list()
            flush_table()
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            flush_para()
            flush_list()
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if not is_table_sep(cells):
                table_rows.append(cells)
            continue
        else:
            flush_table()
        if stripped.startswith("### "):
            flush_para()
            flush_list()
            blocks.append(("h3", stripped[4:].strip()))
        elif stripped.startswith("## "):
            flush_para()
            flush_list()
            blocks.append(("h2", stripped[3:].strip()))
        elif stripped.startswith("# "):
            flush_para()
            flush_list()
            blocks.append(("h1", stripped[2:].strip()))
        elif stripped.startswith(("- ", "* ")):
            flush_para()
            if list_kind == "olist":
                flush_list()
            list_kind = "ulist"
            list_items.append(stripped[2:].strip())
        elif re.match(r"^\d+\.\s", stripped):
            flush_para()
            if list_kind == "ulist":
                flush_list()
            list_kind = "olist"
            list_items.append(re.sub(r"^\d+\.\s", "", stripped))
        else:
            flush_list()
            para_lines.append(stripped)

    flush_para()
    flush_list()
    flush_table()
    if in_code:  # unclosed fence — keep the content anyway
        blocks.append(("code", list(code_lines)))
    return blocks


def inline_segments(text: str) -> list[tuple[str, bool]]:
    """Split ``**bold**`` markup into (text, bold) segments."""
    segments: list[tuple[str, bool]] = []
    pos = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > pos:
            segments.append((text[pos:match.start()], False))
        segments.append((match.group(1), True))
        pos = match.end()
    if pos < len(text):
        segments.append((text[pos:], False))
    return segments or [(text, False)]


# ---------------------------------------------------------------------------
# DOCX backend
# ---------------------------------------------------------------------------

def render_docx(blocks: list[tuple[str, object]], out_path: Path) -> Path:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is not installed (pip install python-docx)") from exc

    doc = Document()
    # Base font for readability.
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    for kind, payload in blocks:
        if kind in ("h1", "h2", "h3"):
            level = {"h1": 1, "h2": 2, "h3": 3}[kind]
            heading = doc.add_heading(level=level)
            for text, bold in inline_segments(str(payload)):
                run = heading.add_run(text)
                run.bold = bold or None
        elif kind == "para":
            para = doc.add_paragraph()
            for text, bold in inline_segments(str(payload)):
                run = para.add_run(text)
                if bold:
                    run.bold = True
        elif kind in ("ulist", "olist"):
            style_name = "List Bullet" if kind == "ulist" else "List Number"
            for item in payload:  # type: ignore[union-attr]
                para = doc.add_paragraph(style=style_name)
                for text, bold in inline_segments(item):
                    run = para.add_run(text)
                    if bold:
                        run.bold = True
        elif kind == "code":
            for code_line in payload:  # type: ignore[union-attr]
                para = doc.add_paragraph()
                run = para.add_run(code_line or " ")
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        elif kind == "table":
            header, rows = payload  # type: ignore[misc]
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Table Grid"
            for j, cell in enumerate(header):
                table.cell(0, j).text = cell
            for i, row in enumerate(rows, start=1):
                for j in range(len(header)):
                    table.cell(i, j).text = row[j] if j < len(row) else ""

    doc.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# PDF backend
# ---------------------------------------------------------------------------

def render_pdf(blocks: list[tuple[str, object]], out_path: Path,
               title: str = "Security Assessment Report") -> Path:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import (ListFlowable, ListItem, Paragraph,
                                        Preformatted, SimpleDocTemplate,
                                        Spacer, Table, TableStyle)
    except ImportError as exc:
        raise RuntimeError(
            "reportlab is not installed (pip install reportlab)") from exc

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("ReportH1", parent=styles["Heading1"],
                              fontSize=20, spaceAfter=12))
    styles.add(ParagraphStyle("ReportH2", parent=styles["Heading2"],
                              fontSize=15, spaceBefore=14, spaceAfter=8))
    styles.add(ParagraphStyle("ReportH3", parent=styles["Heading3"],
                              fontSize=12, spaceBefore=10, spaceAfter=6))
    styles.add(ParagraphStyle("ReportBody", parent=styles["Normal"],
                              fontSize=10, leading=14))
    styles.add(ParagraphStyle("ReportCode", parent=styles["Code"],
                              fontName="Courier", fontSize=8, leading=11))

    def rich(text: str) -> str:
        """**bold** -> <b>, escaped for reportlab's mini-HTML."""
        out = []
        for segment, bold in inline_segments(text):
            seg = html.escape(segment)
            out.append(f"<b>{seg}</b>" if bold else seg)
        return "".join(out)

    story = []
    for kind, payload in blocks:
        if kind == "h1":
            story.append(Paragraph(rich(str(payload)), styles["ReportH1"]))
        elif kind == "h2":
            story.append(Paragraph(rich(str(payload)), styles["ReportH2"]))
        elif kind == "h3":
            story.append(Paragraph(rich(str(payload)), styles["ReportH3"]))
        elif kind == "para":
            story.append(Paragraph(rich(str(payload)), styles["ReportBody"]))
            story.append(Spacer(1, 0.15 * cm))
        elif kind in ("ulist", "olist"):
            items = [ListItem(Paragraph(rich(i), styles["ReportBody"]))
                     for i in payload]  # type: ignore[union-attr]
            story.append(ListFlowable(
                items, bulletType="bullet" if kind == "ulist" else "1",
                leftIndent=0.8 * cm))
            story.append(Spacer(1, 0.15 * cm))
        elif kind == "code":
            for code_line in payload:  # type: ignore[union-attr]
                story.append(Preformatted(html.escape(code_line or " "),
                                          styles["ReportCode"]))
            story.append(Spacer(1, 0.2 * cm))
        elif kind == "table":
            header, rows = payload  # type: ignore[misc]
            data = [[Paragraph(rich(c), styles["ReportBody"]) for c in header]]
            for row in rows:
                data.append([Paragraph(
                    rich(row[j] if j < len(row) else ""), styles["ReportBody"])
                    for j in range(len(header))])
            width = (A4[0] - 4 * cm) / max(len(header), 1)
            table = Table(data, colWidths=[width] * len(header), repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF3")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(table)
            story.append(Spacer(1, 0.2 * cm))

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm,
                               f"Page {_doc.page}")
        canvas.restoreState()

    SimpleDocTemplate(str(out_path), pagesize=A4, title=title,
                      leftMargin=2 * cm, rightMargin=2 * cm).build(
        story, onFirstPage=footer, onLaterPages=footer)
    return out_path


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def convert_report(md_path: Path, fmt: str) -> Path:
    """Convert a Markdown report to ``fmt`` (docx|pdf|md) next to the source.

    ``md`` is a no-op returning the input path.  Raises ValueError for
    unknown formats, RuntimeError when the backend is unavailable.
    """
    fmt = fmt.lower()
    if fmt == "md":
        return md_path
    if fmt not in ("docx", "pdf"):
        raise ValueError(f"unsupported format {fmt!r} (use md, docx, or pdf)")
    if not md_path.exists():
        raise ValueError(f"report not found: {md_path}")

    blocks = parse_markdown(md_path.read_text())
    out_path = md_path.with_suffix(f".{fmt}")
    log.info("Converted %s -> %s", md_path.name, out_path.name)
    if fmt == "docx":
        return render_docx(blocks, out_path)
    return render_pdf(blocks, out_path, title=md_path.stem)
